"""Opt-in real Edge UI/download verification against an offline rebuilt GP snapshot."""
import json
import os
from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.skipif(os.getenv('FINANCIAL_STANDARD_EDGE') != '1', reason='Opt-in rebuilt snapshot Edge verification')
def test_standard_layout_real_edge():
    from playwright.sync_api import sync_playwright
    import openpyxl
    import fitz
    snapshot = json.loads(Path(os.environ['FINANCIAL_REBUILT_SNAPSHOT']).read_text(encoding='utf-8'))
    output = Path(os.environ['FINANCIAL_EDGE_OUTPUT'])
    output.mkdir(parents=True, exist_ok=True)
    errors, receipts = [], []
    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel='msedge', headless=True)
        page = browser.new_page(viewport={'width':1440,'height':1100}, accept_downloads=True)
        page.on('pageerror', lambda error: errors.append(str(error)))
        def route_request(route):
            url = route.request.url
            if not url.startswith('http://financial.local/'):
                route.abort()
                return
            relative = url.split('http://financial.local/', 1)[1] or 'index.html'
            target = ROOT / relative
            if target.is_file():
                route.fulfill(path=str(target), content_type='text/html' if relative.endswith('.html') else 'application/javascript')
            else:
                route.fulfill(status=404, body='Not found')
        page.route('**/*', route_request)
        page.goto('http://financial.local/')
        page.wait_for_function("!document.getElementById('authGate').classList.contains('hidden')")
        scopes = [('SMI',''),('SMI','00'),('SMI','01'),('SMI','02'),('SMI','03'),('RRCOF',''),('BBMTS',''),('SFAB','')]
        for company, branch in scopes:
            payload = next(p['payload'] for p in snapshot['periods'] if p['company_code']==company and p['fiscal_year']==2026 and p['fiscal_period']==6)
            page.evaluate("""payload=>{
                clearReportOnly();state.periodPayload=payload;
                $('companySelect').innerHTML=`<option value="${text(payload.company_code)}">${text(payload.company_name)}</option>`;
                $('monthSelect').innerHTML=`<option data-year="${payload.fiscal_year}" data-period="${payload.fiscal_period}">June</option>`;
                $('companyPickerBtn').textContent=payload.company_name;
                $('yearSelect').innerHTML=`<option value="2026">2026</option>`;
                setStatus('Offline verification — real GP snapshot');
                showApplication({user:{email:'Offline real GP snapshot verification'}});
                populateBranches();applyBranch();selectStatement('income');
            }""",payload)
            if company=='SMI':
                assert page.locator('#branchSelect option').all_text_contents()==['Consolidated — All branches','00 — Fargo','01 — Grand Forks','02 — Fergus Falls','03 — Bismarck']
                if branch:
                    page.select_option('#branchSelect',branch)
            else:
                assert page.locator('#branchFilter').is_hidden()
            name=company+(' branch '+branch if branch else ' consolidated')
            assert page.evaluate('state.report.meta.branch_code')==branch
            assert page.locator('#monthlyRecon').is_disabled()
            page.screenshot(path=str(output/(name+' overview.png')),full_page=False)
            report=page.evaluate("({income:state.report.income,balance:state.report.balance,meta:state.report.meta,totals:state.report.totals,comparison:state.report.comparison,scope_note:state.report.scope_note,ordered:{income:FinancialWorkbook.traditionalRows(state.report.income),balance:FinancialWorkbook.traditionalRows(state.report.balance)}})")
            expected=payload['statement']['branches'][branch] if branch else payload['statement']
            assert report['totals']['net_income']==float(expected['kpis']['net_income'])
            assert page.evaluate('state.charts.trend.data.datasets[3].data')==[float(r['net_income']) for r in expected['trend']]
            for kind,body,tab in [('income','#incomeBody','#incomeTab'),('balance','#balanceBody','#balanceTab')]:
                page.click(tab)
                page.locator('#statements').scroll_into_view_if_needed()
                page.screenshot(path=str(output/(name+' '+kind+'.png')))
                ids=page.locator(body+' tr').evaluate_all('(rows)=>rows.map(r=>r.dataset.rowId)')
                assert ids==[r['id'] for r in report['ordered'][kind]]
                for row in report[kind]:
                    if row['row_type']=='account':
                        assert ids.index(row['id'])<ids.index(row['parent_id'])
                category=next((r['parent_id'] for r in report[kind] if r['row_type']=='account' and r.get('detail')),None)
                if category:
                    page.locator(body+' [data-category="'+category+'"]').click()
                    page.locator('#drilldownBody details').first.locator('summary').click()
                    assert 'raw GP debit / credit' in page.locator('#drilldownBody').inner_text()
            for button,extension in [('#downloadPdf','.pdf'),('#downloadExcel','.xlsx')]:
                with page.expect_download() as pending:
                    page.click(button)
                download=pending.value
                download.save_as(str(output/(name+extension)))
                if branch:
                    assert '-branch-'+branch+'-' in download.suggested_filename
            workbook=openpyxl.load_workbook(output/(name+'.xlsx'))
            assert workbook['Account Detail'].max_row>3 or not any(r.get('detail') or r.get('prior_detail') for kind in ['income','balance'] for r in report[kind])
            assert all(s.sheet_view.pane is None for s in workbook)
            for kind,sheet_name,keys in [('income','Income Statement',['current_period','current_ytd','prior_period','prior_ytd']),('balance','Balance Sheet',['current_balance','prior_balance'])]:
                sheet=workbook[sheet_name]
                for index,row in enumerate(report['ordered'][kind],7):
                    assert sheet.cell(index,1).value.strip()==row['label']
                    for col,key in enumerate(keys,2):
                        value=row.get('values',{}).get(key)
                        assert sheet.cell(index,col).value==(None if value is None else float(value))
            with fitz.open(output/(name+'.pdf')) as pdf:
                text='\n'.join(p.get_text() for p in pdf)
                assert 'Income Statement' in text and 'Balance Sheet' in text
                assert 'June 2026' in text and 'June 2025' in text
                chunks=page.evaluate("['income','balance'].flatMap(type=>pdfPageChunks(pdfRows(state.report[type],type),state.report.scope_note?28:30))")
                assert len(pdf)==len(chunks), 'Unexpected automatic overflow page'
                for i,p in enumerate(pdf):
                    assert len(chunks[i])>=10, 'Orphan statement totals page'
                    assert chunks[i][-1]['row_type']!='heading'
                    p.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(output/(name+f' PDF {i+1}.png'))
                receipts.append({'company':company,'branch':branch,'pdf_pages':len(pdf),'xlsx_rows':{s.title:s.max_row for s in workbook},'row_values_exact':True,'net_income':expected['kpis']['net_income']})
            workbook.close()
        page.set_viewport_size({'width':390,'height':844})
        page.locator('#statements').scroll_into_view_if_needed()
        page.screenshot(path=str(output/'mobile.png'))
        assert not errors, errors
        (output/'edge-receipt.json').write_text(json.dumps({'browser':browser.version,'offline_snapshot':True,'live_auth_tested':False,'scopes':receipts,'page_errors':errors},indent=2),encoding='utf-8')
        browser.close()
