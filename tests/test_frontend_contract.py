"""Static contract tests for the GitHub Pages financial statements client."""
from html.parser import HTMLParser
from pathlib import Path
import json
import re
import subprocess

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INDEX = ROOT / "index.html"


class Markup(HTMLParser):
    def __init__(self):
        super().__init__()
        self.ids: set[str] = set()
        self.scripts: list[dict[str, str | None]] = []
        self.meta: list[dict[str, str | None]] = []
        self.text: list[str] = []

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if values.get("id"):
            self.ids.add(values["id"])
        if tag == "script":
            self.scripts.append(values)
        if tag == "meta":
            self.meta.append(values)

    def handle_data(self, data):
        self.text.append(data)


def source() -> str:
    assert INDEX.exists(), "index.html must exist"
    return INDEX.read_text(encoding="utf-8")


def markup() -> Markup:
    parsed = Markup()
    parsed.feed(source())
    return parsed


def test_single_page_has_required_financial_statement_controls():
    parsed = markup()
    required_ids = {
        "authGate", "password", "loginBtn", "signOutBtn",
        "companySelect", "yearSelect", "monthSelect", "incomeTab", "balanceTab",
        "incomeStatement", "balanceSheet", "accountDrilldown",
        "downloadExcel", "downloadPdf", "trendChart", "positionChart",
        "statusMessage",
    }
    assert required_ids <= parsed.ids
    visible_text = " ".join(parsed.text)
    for label in (
        "Financial Statements", "Income Statement", "Balance Sheet",
        "Download Excel", "Download PDF", "V2.4",
    ):
        assert label in visible_text


def test_client_uses_supabase_password_auth_and_only_required_rpcs():
    html = source()
    assert "window.supabase" in html and "createClient" in html
    assert "signInWithPassword" in html
    assert ".auth.signOut()" in html
    assert "financial_read_current" not in html
    assert re.search(r"\.rpc\(\s*['\"]financial_statement_catalog['\"]\s*\)", html)
    assert re.search(r"\.rpc\(\s*['\"]financial_statement_period['\"]\s*,\s*\{\s*p_company_code\s*:", html)
    for parameter in ("p_company_code", "p_fiscal_year", "p_fiscal_period"):
        assert parameter in html
    assert "smi_sales_current_snapshot" not in html


def test_jon_login_matches_other_apps_pin_flow():
    html = source()
    assert 'label for="password">PIN</label>' in html
    assert "const AUTH_EMAIL='jonj@360-llc.com'" in html
    assert "email:AUTH_EMAIL" in html
    assert 'id="email"' not in html


def test_period_controls_use_separate_year_and_month_and_hide_pre_2020():
    html = source()
    assert 'label for="yearSelect">Year</label>' in html
    assert 'label for="monthSelect">Month</label>' in html
    assert "const MIN_REPORT_YEAR=2020" in html
    assert "Number(fiscalYear(row))>=MIN_REPORT_YEAR" in html
    assert 'id="periodSelect"' not in html


def test_company_picker_avoids_clipped_mobile_native_dropdown():
    html = source()
    assert "COMPANY_PRIORITY" not in html
    assert "String(a[1]).localeCompare(String(b[1]))||String(a[0]).localeCompare(String(b[0]))" in html
    assert 'id="companyPickerBtn"' in html
    assert 'id="companyMenu"' in html
    assert 'role="listbox"' in html
    assert "max-height:min(60vh,420px);overflow-y:auto" in html
    assert "companyMenu').addEventListener('click'" in html
    assert "select.dispatchEvent(new Event('change'))" in html


def test_catalog_rpc_pages_past_supabase_row_limit():
    html = source()
    assert "const CATALOG_PAGE_SIZE=1000" in html
    assert ".range(offset,offset+CATALOG_PAGE_SIZE-1)" in html
    assert "if(page.length<CATALOG_PAGE_SIZE)break" in html
    assert "offset+=CATALOG_PAGE_SIZE" in html


def test_client_fails_closed_without_session_or_when_rpc_fails():
    html = source()
    assert re.search(r"if\s*\(\s*!session(?:\?\.)?(?:access_token)?", html)
    assert "showAuthGate" in html
    assert "clearFinancialData" in html
    assert re.search(r"if\s*\(\s*(?:catalogResult|result|periodResult)\.error\s*\)\s*throw", html)
    assert re.search(r"catch\s*\([^)]*\)\s*\{[^}]*clearFinancialData", html, re.S)
    forbidden = (
        "mockData", "mock_data", "sampleData", "sample_data",
        "demoData", "demo_data", "data/financial", "localhost",
        "127.0.0.1",
    )
    assert not any(token in html for token in forbidden)


def test_financial_payload_is_not_written_to_browser_storage():
    html = source()
    assert "localStorage" not in html
    assert "indexedDB" not in html
    assert "sessionStorage" not in html
    assert re.search(r"persistSession\s*:\s*true", html)
    assert not re.search(r"state\.catalog\s*=.*payload", html)


def test_exports_are_real_pinned_libraries_and_spreadsheet_cells_are_sanitized():
    html = source()
    parsed = markup()
    srcs = [script.get("src") or "" for script in parsed.scripts]
    assert "vendor/exceljs-4.4.0-force-full-calc.min.js" in srcs
    assert "scripts/export_xlsx.js" in srcs
    assert "vendor/jspdf-2.5.2.min.js" in srcs
    assert "vendor/jspdf-autotable-3.8.4.min.js" in srcs
    assert all(not re.match(r"https?://", src) for src in srcs)
    for src in srcs:
        if src:
            assert (ROOT / src).is_file(), f"vendored script is missing: {src}"
    assert "workbook.xlsx.writeBuffer" in html
    assert "jsPDF" in html and "autoTable" in html and ".save(" in html
    assert "sanitizeSpreadsheetCell" in html
    assert re.search(r"^[^\n]*[=+\-@]", html, re.M), "formula-leading characters must be handled"


def test_charts_and_drilldown_are_data_driven():
    html = source()
    assert "vendor/chart-4.4.7.min.js" in html
    for metric in (
        "revenue", "gross_profit", "expenses", "net_income",
        "assets", "liabilities", "equity",
    ):
        assert metric in html
    assert "renderTrendChart" in html
    assert "renderPositionChart" in html
    assert "renderDrilldown" in html
    assert "row.parent_id===category" in html
    assert "row.values?.[valueKey]" in html
    assert "category-row" in html and "data-category" in html
    assert ".drill-panel{display:block}" in html
    assert "Choose a statement section" in html
    assert re.search(r"addEventListener\(\s*['\"]click['\"]", html)


def test_client_consumes_exact_nested_backend_payload_without_duplicating_raw_accounts():
    html = source()
    assert "root.payload??root" in html
    assert "nested.statement" in html
    assert "statement.presentation" in html
    assert "statement.kpis" in html
    assert "statement.trend" in html
    assert "rawIncome" not in html
    assert "rawBalance" not in html
    assert "normalizeAccount" not in html
    assert "fiscal_period" in html
    assert "period_id" not in html
    assert "root.payload??root" in html
    assert "Net Income" in html
    assert "Total Assets" in html
    assert "Total Liabilities &amp; Equity" in html
    assert "Statement total" not in html


def test_lender_view_uses_backend_presentation_with_comparative_columns():
    html = source()
    assert "statement.presentation" in html
    assert "presentation.income_statement" in html
    assert "presentation.balance_sheet" in html
    for field in (
        "current_period", "current_ytd", "prior_period", "prior_ytd",
        "current_balance", "prior_balance",
    ):
        assert field in html
    assert 'id="incomePeriodHead"' in html
    assert 'id="incomeYtdHead"' in html
    assert 'id="incomePriorPeriodHead"' in html
    assert 'id="incomePriorYtdHead"' in html
    assert 'id="balanceCurrentHead"' in html
    assert 'id="balancePriorHead"' in html
    assert 'class="statement-report-header"' in html
    assert "row.row_type" in html and "row.level" in html
    assert "function strictNumber" in html
    assert "new Set" in html
    assert "Presentation tie-out failed" in html
    assert "Number(value)||0" not in html


def test_excel_export_is_comparative_structured_and_lender_ready():
    html = source()
    exporter = (ROOT / "scripts" / "export_xlsx.js").read_text(encoding="utf-8")
    assert "vendor/exceljs-4.4.0-force-full-calc.min.js" in html
    assert "scripts/export_xlsx.js" in html
    assert "FinancialWorkbook.buildFinancialWorkbook" in html
    assert "workbook.xlsx.writeBuffer" in html
    assert '"$"#,##0;("$"#,##0);-' in exporter
    assert "current_period" in exporter and "prior_ytd" in exporter
    assert "current_balance" in exporter and "prior_balance" in exporter
    assert "source.account_number ?" not in exporter
    assert "row.account_number?`${row.account_number}  `" not in html


def test_exceljs_generated_file_preserves_lender_formatting(tmp_path):
    report = {
        "meta": {"company_name": "Structural Fab LLC", "fiscal_year": 2026, "fiscal_period": 5},
        "comparison": {"fiscal_year": 2025, "available": True},
        "income": [
            {"id": "revenue", "row_type": "section", "level": 0, "label": "Revenue", "values": {"current_period": "100", "current_ytd": "500", "prior_period": "90", "prior_ytd": "450"}},
            {"id": "net_income", "row_type": "subtotal", "level": 0, "label": "Net Income", "values": {"current_period": "34", "current_ytd": "170", "prior_period": "30", "prior_ytd": "150"}},
        ],
        "balance": [
            {"id": "assets", "row_type": "heading", "level": 0, "label": "Assets", "values": {"current_balance": "100", "prior_balance": "90"}},
            {"id": "total_assets", "row_type": "subtotal", "level": 0, "label": "Total Assets", "values": {"current_balance": "100", "prior_balance": "90"}},
            {"id": "total_liabilities_and_equity", "row_type": "subtotal", "level": 0, "label": "Total Liabilities & Equity", "values": {"current_balance": "100", "prior_balance": "90"}},
        ],
    }
    output = tmp_path / "lender.xlsx"
    runner = tmp_path / "build.js"
    runner.write_text(
        "const fs=require('fs');const api=require(process.argv[2]);"
        "const report=JSON.parse(fs.readFileSync(process.argv[3],'utf8'));"
        "(async()=>{const wb=api.buildFinancialWorkbook(report,{month:'May'});"
        "const data=await wb.xlsx.writeBuffer();fs.writeFileSync(process.argv[4],Buffer.from(data));})()"
        ".catch(error=>{console.error(error);process.exit(1)});",
        encoding="utf-8",
    )
    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps(report), encoding="utf-8")
    subprocess.run(
        ["node", str(runner), str(ROOT / "scripts" / "export_xlsx.js"), str(report_path), str(output)],
        check=True,
    )
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["Income Statement", "Balance Sheet", "Controls"]
    income = workbook["Income Statement"]
    assert "A1:E1" in {str(region) for region in income.merged_cells.ranges}
    assert income["A1"].value == "Structural Fab LLC"
    assert income["A1"].font.bold and income["A1"].font.sz == 16
    assert income["A6"].fill.fgColor.rgb.endswith("233B61")
    assert income["A8"].border.bottom.style == "double"
    assert income["B7"].value == 100 and income["B7"].number_format == '"$"#,##0;("$"#,##0);-'
    assert income.column_dimensions["A"].width >= 40
    assert income.freeze_panes == "B7"
    assert income.sheet_properties.pageSetUpPr.fitToPage
    assert income.page_setup.fitToHeight == 1
    assert income.page_setup.orientation == "landscape" and income.page_setup.fitToWidth == 1
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    controls = workbook["Controls"]
    assert controls.sheet_state == "hidden"
    assert controls["B2"].data_type == "f" and "'Balance Sheet'!B" in controls["B2"].value


def test_pdf_export_has_lender_headers_repeating_columns_and_page_numbers():
    html = source()
    assert "orientation:'landscape'" in html
    assert "showHead:'everyPage'" in html
    assert "Accrual Basis" in html
    assert "Unaudited" in html
    assert "Page ${page} of ${pages}" in html
    assert "putTotalPages" not in html
    assert "rowPageBreak:'avoid'" in html
    assert "PDF export failed" in html and "try{setBusy(true)" in html
    assert "row.raw.row_type" in html


def test_page_is_responsive_accessible_and_blue_white_without_purple():
    html = source()
    parsed = markup()
    assert any(meta.get("name") == "viewport" for meta in parsed.meta)
    assert "@media" in html
    assert ".main{max-width:100vw;overflow-x:hidden}" in html
    assert ".top-row{align-items:flex-start;flex-wrap:wrap}" in html
    assert "aria-live" in html
    assert ":focus-visible" in html
    assert "--blue" in html and "#ffffff" in html.lower()
    forbidden_colors = ("purple", "violet", "#800080", "#7c3aed", "#8b5cf6", "#a855f7")
    assert not any(color in html.lower() for color in forbidden_colors)


def test_authenticated_page_has_no_third_party_runtime_scripts_and_uses_csp():
    html = source()
    parsed = markup()
    assert 'http-equiv="Content-Security-Policy"' in html
    assert "script-src 'self' 'unsafe-inline'" in html
    assert "connect-src 'self' https://inhwadbibwkakacdvoxu.supabase.co" in html
    assert "cdn.jsdelivr.net" not in html
    assert "cdn.sheetjs.com" not in html
    assert all(not re.match(r"https?://", script.get("src") or "") for script in parsed.scripts)
    assert "vendor/supabase-2.49.1.min.js" in html
